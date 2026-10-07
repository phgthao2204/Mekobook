from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\Mekobook")
OUTPUT = ROOT / "docs" / "BAO_CAO_NHAN_DIEN_THUONG_HIEU_MEKOBOOK.docx"
LOGIN_SCREEN = Path(r"C:\Users\Dell\AppData\Local\Temp\codex-clipboard-5bd61c1c-5917-4676-aed5-aeb67397860f.png")
LIBRARY_SCREEN = Path(r"C:\Users\Dell\AppData\Local\Temp\codex-clipboard-886c54f8-ffca-4f26-9a89-9383c664ab12.png")
APP_ICON = ROOT / "assets" / "icon.png"

EMERALD = "059669"
NAVY = "0F172A"
SLATE = "64748B"
LIGHT_BORDER = "D9D9D9"
PALE_GREEN = "ECFDF5"
PALE_BLUE = "F1F5F9"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(SLATE)
    return p


def add_bullet(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(NAVY)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.15

title_style = styles["Title"]
title_style.font.name = "Arial"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
title_style.font.size = Pt(26)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(10)

for style_name, size in (("Heading 1", 18), ("Heading 2", 13)):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.keep_with_next = True

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_p.add_run("Mekobook   |   Báo cáo nhận diện thương hiệu   |   ").font.size = Pt(8)
add_page_number(footer_p)
for run in footer_p.runs:
    run.font.color.rgb = RGBColor.from_string(SLATE)

# Cover
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(24)
badge = p.add_run("▣")
badge.font.size = Pt(40)
badge.font.color.rgb = RGBColor.from_string(EMERALD)

title = doc.add_paragraph("Báo cáo nhận diện thương hiệu Mekobook", style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph("Tổng hợp từ giao diện và mã nguồn ứng dụng di động")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.runs[0].font.size = Pt(13)
subtitle.runs[0].font.color.rgb = RGBColor.from_string(SLATE)

if LOGIN_SCREEN.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.add_run().add_picture(str(LOGIN_SCREEN), width=Inches(2.25))
    add_caption(doc, "Hình 1. Màn hình đăng nhập thể hiện ngôn ngữ nhận diện hiện tại")

intro = doc.add_paragraph()
intro.paragraph_format.space_before = Pt(5)
intro.add_run("Kết luận chính. ").bold = True
intro.add_run(
    "Mekobook đã hình thành một nhận diện rõ nét với màu xanh emerald, biểu tượng sách mở, "
    "kiểu chữ đậm và hệ giao diện bo tròn. Hệ thống này phù hợp với một thư viện số hiện đại, "
    "đáng tin cậy và thân thiện. Điểm cần ưu tiên là đồng bộ icon cài đặt và splash screen với "
    "logo sách đang xuất hiện bên trong ứng dụng."
)

doc.add_page_break()

# 1. Color and typography
doc.add_heading("1 Màu sắc và font chữ", level=1)
doc.add_paragraph(
    "Bảng màu trong mã nguồn lấy xanh emerald làm màu nhận diện chính, kết hợp nền sáng và các "
    "sắc slate trung tính. Cách phối này tạo cảm giác sạch, bình tĩnh và đủ tương phản cho sản phẩm đọc sách."
)

rows = [
    ("Thương hiệu", "#059669", "Nút chính, logo chữ, trạng thái chủ động", EMERALD),
    ("Nền ứng dụng", "#F8FAFC", "Nền tổng thể", "F8FAFC"),
    ("Bề mặt", "#FFFFFF", "Thẻ, form, modal", WHITE),
    ("Chữ chính", "#0F172A", "Tiêu đề và nội dung quan trọng", NAVY),
    ("Chữ phụ", "#64748B", "Mô tả, metadata, trạng thái phụ", SLATE),
    ("Đường viền", "#E2E8F0", "Phân tách nhẹ giữa các bề mặt", "E2E8F0"),
    ("Lỗi", "#B91C1C", "Thông báo lỗi và hành động nguy hiểm", "B91C1C"),
]
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Cm(3.2), Cm(2.5), Cm(8.0), Cm(1.5)]
headers = ["Vai trò", "Mã màu", "Ứng dụng", "Mẫu"]
for i, text in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.width = widths[i]
    cell.text = text
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for run in cell.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
        run.font.size = Pt(9.5)
for idx, row in enumerate(rows):
    cells = table.add_row().cells
    for i, value in enumerate(row[:3]):
        cells[i].width = widths[i]
        cells[i].text = value
        set_cell_margins(cells[i])
        cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if idx % 2:
            set_cell_shading(cells[i], "F8FAFC")
    cells[3].width = widths[3]
    cells[3].text = ""
    set_cell_shading(cells[3], row[3])
set_table_borders(table)

doc.add_heading("Hệ màu hỗ trợ", level=2)
add_bullet(doc, "Xanh nhạt #ECFDF5 và xanh đậm #065F46 dành cho badge người dùng và trạng thái tích cực.")
add_bullet(doc, "Xanh dương #3B82F6 biểu thị tiến độ đọc, tránh cạnh tranh với màu thương hiệu chính.")
add_bullet(doc, "Vàng #B45309 và #FEF3C7 dành cho bản quyền hoặc cảnh báo cần chú ý.")
add_bullet(doc, "Đỏ nhạt #FEF2F2 kết hợp đỏ đậm dành cho đăng xuất và lỗi.")

doc.add_heading("Typography", level=2)
doc.add_paragraph(
    "Ứng dụng chưa nạp font tùy chỉnh và không khai báo fontFamily. Vì vậy Android sử dụng Roboto, "
    "iOS sử dụng San Francisco và web dùng font hệ thống. Wordmark MEKOBOOK được viết hoa, độ đậm 900, "
    "giãn chữ 1,5 đến 2 điểm. Tiêu đề dùng độ đậm 700 đến 800; nội dung dùng 400 đến 600."
)
doc.add_paragraph(
    "Khuyến nghị: tiếp tục dùng system font để giữ hiệu năng, hoặc chọn Be Vietnam Pro làm font thương hiệu "
    "nếu cần nhận diện riêng và khả năng hiển thị tiếng Việt nhất quán trên nhiều nền tảng."
)

doc.add_page_break()

# 2. Logo and imagery
doc.add_heading("2 Logo và hình ảnh", level=1)
doc.add_paragraph(
    "Logo đang xuất hiện trong giao diện gồm biểu tượng quyển sách mở màu trắng đặt trong khối vuông bo góc "
    "màu xanh emerald, đi cùng chữ MEKOBOOK viết hoa. Biểu tượng truyền tải trực tiếp lĩnh vực đọc sách và tri thức; "
    "khối bo góc giúp hình ảnh phù hợp với môi trường ứng dụng di động."
)

image_table = doc.add_table(rows=1, cols=2)
image_table.alignment = WD_TABLE_ALIGNMENT.CENTER
image_table.autofit = False
for cell in image_table.rows[0].cells:
    cell.width = Cm(7.5)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell, top=160, bottom=160, start=160, end=160)
set_table_borders(image_table)

if LIBRARY_SCREEN.exists():
    p = image_table.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(LIBRARY_SCREEN), width=Inches(2.15))
    cp = image_table.cell(0, 0).add_paragraph("Giao diện sử dụng xanh emerald và icon dạng nét")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.runs[0].italic = True
    cp.runs[0].font.size = Pt(8.5)

if APP_ICON.exists():
    p = image_table.cell(0, 1).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(APP_ICON), width=Inches(2.15))
    cp = image_table.cell(0, 1).add_paragraph("Icon cài đặt hiện vẫn là asset mẫu của Expo")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.runs[0].italic = True
    cp.runs[0].font.size = Pt(8.5)

doc.add_paragraph(
    "Khoảng cách giữa nhận diện trong ứng dụng và icon cài đặt là vấn đề lớn nhất hiện nay. Người dùng nhìn thấy "
    "biểu tượng sách trong màn hình đăng nhập nhưng lại thấy biểu tượng Expo trên màn hình chính của điện thoại. "
    "Điều này làm giảm khả năng ghi nhớ và tạo cảm giác sản phẩm chưa hoàn thiện."
)

doc.add_heading("Bộ logo cần chuẩn hóa", level=2)
add_bullet(doc, "Logo đầy đủ gồm biểu tượng sách và chữ MEKOBOOK.")
add_bullet(doc, "Logo rút gọn chỉ gồm biểu tượng sách để dùng ở kích thước nhỏ.")
add_bullet(doc, "App icon Android và iOS với vùng an toàn phù hợp từng nền tảng.")
add_bullet(doc, "Adaptive icon gồm foreground và background tách biệt cho Android.")
add_bullet(doc, "Splash icon, favicon web và phiên bản đơn sắc.")

doc.add_heading("Nguyên tắc sử dụng hình ảnh", level=2)
doc.add_paragraph(
    "Ảnh bìa sách là hình ảnh trung tâm của sản phẩm. Bìa cần giữ đúng tỷ lệ dọc, không phủ bộ lọc làm thay đổi "
    "thiết kế gốc và luôn có placeholder mang nhận diện Mekobook khi API không cung cấp ảnh. Nền xung quanh nên "
    "trung tính để ưu tiên nội dung sách."
)

doc.add_page_break()

# 3. Shape language
doc.add_heading("3 Ngôn ngữ hình khối và biểu tượng", level=1)
doc.add_paragraph(
    "Ngôn ngữ hình khối là tập hợp quy tắc về bo góc, đường viền, icon và cách sắp xếp thành phần. Đây là mục "
    "giúp các màn hình mới vẫn trông như cùng một sản phẩm, ngay cả khi do nhiều người thiết kế hoặc phát triển."
)

shape_table = doc.add_table(rows=1, cols=3)
shape_table.alignment = WD_TABLE_ALIGNMENT.CENTER
shape_table.autofit = False
shape_headers = ["Thành phần", "Quy cách hiện tại", "Ý nghĩa sử dụng"]
for i, text in enumerate(shape_headers):
    cell = shape_table.rows[0].cells[i]
    cell.text = text
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell)
    for run in cell.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
        run.font.size = Pt(9.5)
shape_rows = [
    ("Badge và nút nhỏ", "Bo góc 6 đến 11 px", "Trạng thái, hành động phụ, input"),
    ("Thanh tìm kiếm và thẻ", "Bo góc 12 đến 18 px", "Bề mặt chính và nhóm nội dung"),
    ("Modal hồ sơ", "Bo góc 20 px", "Tạo cảm giác mềm và thân thiện"),
    ("Badge người dùng", "Dạng viên thuốc", "Danh tính và trạng thái trực tuyến"),
    ("Đường viền", "1 px, màu xám nhạt", "Phân lớp nhẹ, tránh cảm giác nặng"),
]
for idx, row in enumerate(shape_rows):
    cells = shape_table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
        set_cell_margins(cells[i])
        cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if idx % 2:
            set_cell_shading(cells[i], "F8FAFC")
set_table_borders(shape_table)

doc.add_heading("Hệ biểu tượng", level=2)
doc.add_paragraph(
    "Ứng dụng dùng Ionicons làm bộ icon duy nhất. Icon quyển sách đại diện cho thương hiệu và hành động đọc; "
    "person đại diện tài khoản hoặc tác giả; search dành cho tìm kiếm; shield-checkmark và lock-open thể hiện "
    "quyền truy cập; grid và list chuyển chế độ hiển thị. Phong cách chính là icon outline, trong khi trạng thái "
    "được chọn có thể dùng icon đặc."
)
add_bullet(doc, "Không trộn nhiều bộ icon trong cùng một màn hình.")
add_bullet(doc, "Dùng cùng độ dày nét và kích thước quang học cho các icon cùng cấp.")
add_bullet(doc, "Màu icon chủ động dùng xanh emerald; icon phụ dùng slate; lỗi dùng đỏ.")

doc.add_heading("Khoảng cách và nhịp điệu", level=2)
doc.add_paragraph(
    "Giao diện hiện dùng khoảng trắng rộng, card tách biệt và độ cao nút khoảng 46 đến 50 px. Đây là nền tảng tốt "
    "cho khả năng chạm trên điện thoại. Nên chuẩn hóa spacing theo thang 4, 8, 12, 16, 24 và 32 px để giảm "
    "các giá trị rời rạc và giúp phát triển màn hình mới nhanh hơn."
)

doc.add_page_break()

# 4. Tone and experience
doc.add_heading("4 Giọng điệu thương hiệu và trải nghiệm", level=1)
doc.add_paragraph(
    "Giọng điệu thương hiệu quy định cách Mekobook nói chuyện với người đọc trong tiêu đề, nút, trạng thái và "
    "thông báo lỗi. Đây không chỉ là nội dung chữ; nó quyết định cảm giác tin cậy và mức độ dễ sử dụng của sản phẩm."
)

traits = doc.add_table(rows=1, cols=2)
traits.alignment = WD_TABLE_ALIGNMENT.CENTER
traits.autofit = False
for i, text in enumerate(("Đặc điểm", "Biểu hiện trong sản phẩm")):
    cell = traits.rows[0].cells[i]
    cell.text = text
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell)
    for run in cell.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
        run.font.size = Pt(9.5)
trait_rows = [
    ("Tin cậy", "Màu emerald, trạng thái bản quyền và thông tin phiên đọc rõ ràng."),
    ("Hiện đại", "Bố cục tối giản, card, badge và khoảng trắng có chủ đích."),
    ("Thân thiện", "Góc bo tròn, câu chữ gần gũi và hành động dễ hiểu."),
    ("Tập trung vào việc đọc", "Bìa sách, tiến độ và nút Đọc tiếp được ưu tiên thị giác."),
]
for idx, row in enumerate(trait_rows):
    cells = traits.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
        set_cell_margins(cells[i])
        if idx % 2:
            set_cell_shading(cells[i], "F8FAFC")
set_table_borders(traits)

doc.add_heading("Nguyên tắc viết nội dung", level=2)
add_bullet(doc, "Viết ngắn gọn và dùng động từ rõ ràng, ví dụ Đọc tiếp, Đăng nhập, Thử lại.")
add_bullet(doc, "Nói rõ người dùng cần làm gì khi gặp lỗi.")
add_bullet(doc, "Không hiển thị thuật ngữ kỹ thuật như OAuth, token hoặc HTTP 403 cho người đọc phổ thông.")
add_bullet(doc, "Dùng tiếng Việt thống nhất, tự nhiên và tránh pha trộn tiếng Anh khi không cần thiết.")
add_bullet(doc, "Phân biệt cảnh báo, lỗi và trạng thái thông tin bằng cả màu sắc lẫn câu chữ.")

doc.add_heading("Ví dụ diễn đạt", level=2)
examples = doc.add_table(rows=1, cols=2)
examples.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, text in enumerate(("Nên dùng", "Không nên dùng")):
    cell = examples.rows[0].cells[i]
    cell.text = text
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell)
    for run in cell.paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
example_rows = [
    ("Không thể đăng nhập. Vui lòng kiểm tra tài khoản và thử lại.", "invalid_grant hoặc HTTP 401"),
    ("Không thể tải hồ sơ. Thư viện vẫn có thể sử dụng.", "Profile endpoint returned 403"),
    ("Mất kết nối với máy chủ. Chạm để thử lại.", "Network request failed"),
]
for idx, row in enumerate(example_rows):
    cells = examples.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = value
        set_cell_margins(cells[i])
        if idx % 2:
            set_cell_shading(cells[i], "F8FAFC")
set_table_borders(examples)

doc.add_page_break()

# 5. Recommendations
doc.add_heading("5 Đề xuất chuẩn hóa", level=1)
doc.add_paragraph(
    "Để chuyển nhận diện hiện tại thành một hệ thống có thể sử dụng lâu dài, nhóm phát triển nên ưu tiên các công việc sau."
)

recommendations = [
    ("Ưu tiên 1", "Thiết kế và thay thế bộ app icon, adaptive icon, splash icon và favicon bằng biểu tượng sách Mekobook."),
    ("Ưu tiên 2", "Đưa toàn bộ màu semantic, radius, spacing và typography token về một tệp theme dùng chung."),
    ("Ưu tiên 3", "Tạo logo master dạng vector và quy định vùng an toàn, kích thước tối thiểu, nền được phép sử dụng."),
    ("Ưu tiên 4", "Chuẩn hóa nội dung thông báo lỗi theo ngôn ngữ người dùng, không để lộ thuật ngữ hạ tầng."),
    ("Ưu tiên 5", "Bổ sung placeholder bìa sách và trạng thái empty/loading theo cùng hệ nhận diện."),
]
for lead, body in recommendations:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(f"{lead}. ")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(EMERALD)
    p.add_run(body)

doc.add_heading("Tuyên bố nhận diện đề xuất", level=2)
doc.add_paragraph(
    "Mekobook là thư viện số hiện đại dành cho trải nghiệm đọc tập trung và đáng tin cậy. Nhận diện sử dụng xanh "
    "emerald làm dấu hiệu chính, biểu tượng sách mở làm tài sản thị giác trung tâm, typography rõ ràng và hệ giao diện "
    "mềm mại để người đọc tiếp cận nội dung nhanh chóng trên mọi thiết bị."
)

doc.add_heading("Nguồn đối chiếu", level=2)
source_p = doc.add_paragraph()
source_p.add_run("Mã nguồn: ").bold = True
source_p.add_run("src/constants/theme.ts, src/screens/auth/LoginScreen.tsx, src/screens/books/BookCatalog.tsx, app.json và thư mục assets.")
source_p = doc.add_paragraph()
source_p.add_run("Hình minh họa: ").bold = True
source_p.add_run("ảnh chụp ứng dụng do người dùng cung cấp và icon hiện tại trong assets/icon.png.")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
